from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.models import Student
from tests.conftest import authed_client


def _finalize_one_session(client, base_data):
    s1, s2, s3 = [student.id for student in base_data.students[:3]]
    response = client.post(
        "/api/v1/attendance/sessions",
        json={
            "class_id": base_data.class_one.id,
            "subject_id": base_data.subject_one.id,
            "session_date": "2026-10-05",
            "session_number": 1,
        },
    )
    session_id = response.json()["id"]
    client.put(
        f"/api/v1/attendance/sessions/{session_id}/records",
        json={
            "records": [
                {"student_id": s1, "status": "present"},
                {"student_id": s2, "status": "absent"},
                {"student_id": s3, "status": "present"},
            ]
        },
    )
    client.patch(
        f"/api/v1/attendance/sessions/{session_id}/finalize",
        json={"is_finalized": True},
    )
    return session_id


def test_session_csv_export_contains_register_and_counts(client, base_data):
    authed_client(client, "faculty@test.edu")
    session_id = _finalize_one_session(client, base_data)
    response = client.get(f"/api/v1/attendance/sessions/{session_id}/export?format=csv")
    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    text = response.content.decode("utf-8-sig")
    assert "Present: 2" in text
    assert "Absent: 1" in text
    assert "Aarav Menon" in text
    assert "Finalized" in text
    disposition = response.headers["content-disposition"]
    assert "attendance_iii_cse_2026_10_05" in disposition


def test_attendance_summary_export_csv_and_xlsx(client, base_data):
    authed_client(client, "faculty@test.edu")
    _finalize_one_session(client, base_data)
    params = {"report": "students", "class_id": base_data.class_one.id, "format": "csv"}
    response = client.get("/api/v1/attendance/reports/export", params=params)
    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    flattened = "\n".join(",".join(row) for row in rows)
    assert "Student-wise Attendance Summary" in flattened
    assert "Aarav Menon" in flattened
    assert "Percentage" in flattened

    params["format"] = "xlsx"
    xlsx = client.get("/api/v1/attendance/reports/export", params=params)
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    sheet = workbook.active
    values = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
    assert any("Aarav Menon" in value for value in values)


def test_formula_injection_is_neutralised_in_csv_and_xlsx(client, base_data):
    authed_client(client, "faculty@test.edu")
    s1 = base_data.students[0]
    s1.full_name = "=HYPERLINK(\"http://evil.test\",\"click\")"
    s1.roll_number = "+CS-0001"
    with SessionLocal() as db:
        student = db.get(Student, s1.id)
        student.full_name = s1.full_name
        student.roll_number = s1.roll_number
        db.commit()

    _finalize_one_session(client, base_data)
    response = client.get(
        "/api/v1/attendance/reports/export",
        params={"report": "students", "class_id": base_data.class_one.id, "format": "csv"},
    )
    text = response.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text
    assert "\n=HYPERLINK" not in text
    assert "'+CS-0001" in text

    xlsx = client.get(
        "/api/v1/attendance/reports/export",
        params={"report": "students", "class_id": base_data.class_one.id, "format": "xlsx"},
    )
    workbook = load_workbook(io.BytesIO(xlsx.content))
    sheet = workbook.active
    cell_values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert any(isinstance(value, str) and value.startswith("'=HYPERLINK") for value in cell_values)
    assert not any(isinstance(value, str) and value.startswith("=") for value in cell_values)


def test_history_export_and_invalid_format(client, base_data):
    authed_client(client, "faculty@test.edu")
    _finalize_one_session(client, base_data)
    history = client.get(
        "/api/v1/attendance/reports/export?report=history&format=csv"
    )
    assert history.status_code == 200
    text = history.content.decode("utf-8-sig")
    assert "Attendance History" in text
    assert "Rohan D'Souza" in text

    invalid = client.get(
        "/api/v1/attendance/reports/export?report=history&format=pdf"
    )
    assert invalid.status_code == 400
    unknown = client.get("/api/v1/attendance/reports/export?report=magic")
    assert unknown.status_code == 400


def test_student_cannot_export_other_reports(client, base_data):
    authed_client(client, "student@test.edu")
    assert (
        client.get("/api/v1/attendance/reports/export?report=history").status_code == 403
    )
    own = client.get(
        "/api/v1/attendance/reports/export",
        params={"report": "student", "format": "csv"},
    )
    assert own.status_code == 200
    assert "Aarav Menon" in own.content.decode("utf-8-sig")
