"""CSV / XLSX export helpers with Unicode and formula-injection protection."""

from __future__ import annotations

import codecs
import csv
import io
import re
from datetime import date, datetime, timezone
from typing import Any, Iterable, Sequence
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.responses import StreamingResponse

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: Any) -> Any:
    """Neutralise spreadsheet formula injection for untrusted text values."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def slugify(value: str, fallback: str = "export") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_").lower()
    return cleaned[:80] or fallback


def build_filename(name: str, extension: str) -> str:
    stem = slugify(name)
    return f"{stem}.{extension}"


def _attachment_headers(filename: str) -> dict[str, str]:
    fallback = re.sub(r"[^A-Za-z0-9._-]", "_", filename)
    return {
        "Content-Disposition": (
            f'attachment; filename="{fallback}"; filename*=UTF-8\'\'{quote(filename)}'
        )
    }


def _format_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return safe_text(value)


def _excel_value(value: Any) -> Any:
    """openpyxl rejects timezone-aware datetimes, so store UTC naive values."""
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(timezone.utc).replace(tzinfo=None)
        return value
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return safe_text(value)
    return value


def csv_response(
    rows: Iterable[dict[str, Any]],
    headers: Sequence[tuple[str, str]],
    filename: str,
    title: str | None = None,
    summary: Sequence[str] = (),
) -> StreamingResponse:
    buffer = io.StringIO()
    buffer.write(codecs.BOM_UTF8.decode("utf-8"))
    writer = csv.writer(buffer, lineterminator="\r\n")
    if title:
        writer.writerow([safe_text(title)])
        writer.writerow([])
    for line in summary:
        writer.writerow([safe_text(line)])
    if summary or title:
        writer.writerow([])
    writer.writerow([label for label, _ in headers])
    for row in rows:
        writer.writerow([_format_value(row.get(key)) for _, key in headers])
    content = buffer.getvalue().encode("utf-8")
    return StreamingResponse(
        iter([content]),
        media_type="text/csv; charset=utf-8",
        headers=_attachment_headers(filename),
    )


def xlsx_response(
    rows: Iterable[dict[str, Any]],
    headers: Sequence[tuple[str, str]],
    filename: str,
    title: str | None = None,
    summary: Sequence[str] = (),
) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "_", (title or "Export"))[:31] or "Export"

    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="1E293B")
    summary_font = Font(bold=False, size=10, color="334155")

    current_row = 1
    if title:
        sheet.cell(row=current_row, column=1, value=safe_text(title))
        sheet.cell(row=current_row, column=1).font = title_font
        sheet.merge_cells(
            start_row=current_row, start_column=1, end_row=current_row, end_column=max(len(headers), 1)
        )
        current_row += 1
    for line in summary:
        sheet.cell(row=current_row, column=1, value=safe_text(line))
        sheet.cell(row=current_row, column=1).font = summary_font
        sheet.merge_cells(
            start_row=current_row, start_column=1, end_row=current_row, end_column=max(len(headers), 1)
        )
        current_row += 1
    if title or summary:
        current_row += 1

    header_row = current_row
    for col_index, (label, _) in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    current_row += 1

    widths = [len(label) for label, _ in headers]
    for row in rows:
        for col_index, (_, key) in enumerate(headers, start=1):
            raw_value = row.get(key)
            value = _excel_value(raw_value)
            cell = sheet.cell(row=current_row, column=col_index, value=value)
            if isinstance(raw_value, (datetime, date)):
                cell.number_format = "yyyy-mm-dd hh:mm" if isinstance(raw_value, datetime) else "yyyy-mm-dd"
            widths[col_index - 1] = max(widths[col_index - 1], min(len(str(value or "")), 40))
        current_row += 1

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 42)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers=_attachment_headers(filename),
    )


def export_response(
    rows: Iterable[dict[str, Any]],
    headers: Sequence[tuple[str, str]],
    filename_stem: str,
    export_format: str,
    title: str | None = None,
    summary: Sequence[str] = (),
) -> StreamingResponse:
    fmt = export_format.lower()
    if fmt not in ("csv", "xlsx"):
        raise ValueError("Unsupported export format")
    filename = build_filename(filename_stem, fmt)
    if fmt == "csv":
        return csv_response(rows, headers, filename, title=title, summary=summary)
    return xlsx_response(rows, headers, filename, title=title, summary=summary)
